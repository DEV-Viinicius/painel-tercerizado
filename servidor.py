"""
Servidor local do PAINEL DE ATENDIMENTOS.

Cada PC roda o seu proprio painel (127.0.0.1:8000) e todos gravam no MESMO
dados.json da pasta de rede, protegido por uma trava de arquivo.

Inicie com um duplo clique no painel.exe (ou: python servidor.py).

OBS: reconstruido a partir do bytecode do painel.exe (o fonte original foi
apagado da pasta TERCERIZADO). A pagina HTML vive agora em painel.html e e
transformada no modulo pagina.py por gerar_pagina.py antes do build.
"""

import os
import io
import re
import time
import json
import base64
import socket
import threading
import webbrowser
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

import openpyxl
import xlrd

from unificar_atendimentos import ler_terceirizados, ler_equipamentos, limpa, norm
from pagina import PAGINA
import chamados as _ch


def _resolver_pasta_dados():
    """Acha a pasta de rede onde moram o dados.json e as planilhas."""
    for c in (os.environ.get('PAINEL_DADOS'),
              r'\\192.168.0.5\Tecnologia\DOCUMENTACOES E TUTORIAIS\Projetos de Setores\Tercerizado',
              r'Y:\DOCUMENTACOES E TUTORIAIS\Projetos de Setores\Tercerizado'):
        if c and os.path.isdir(c):
            return c
    # Sem rede: usa a pasta do proprio programa (dados ficam SO neste PC).
    return os.path.dirname(os.path.abspath(__file__))


PASTA = _resolver_pasta_dados()
ARQ_DADOS = os.path.join(PASTA, 'dados.json')
ARQ_LOGO = os.path.join(PASTA, 'LOGO SOLIVETTI.jpg')
ARQ_TERC = os.path.join(PASTA, 'PLANILHA APOIO - ATENDIMENTOS TERCEIRIZADOS (1).xlsx')
ARQ_EQUIP_XLSX = os.path.join(PASTA, 'Relatório Equipamentos - Completo.xlsx')
# Porta fixa 8000 (todo mundo acessa localhost:8000). PAINEL_PORTA so e usada
# para testar uma versao nova sem derrubar o painel que ja esta aberto.
try:
    PORTA = int(os.environ.get('PAINEL_PORTA') or 8000)
except ValueError:
    PORTA = 8000

import unificar_atendimentos as _u
_u.PASTA = PASTA
_u.ARQ_TERC = ARQ_TERC
_u.ARQ_EQUIP = os.path.join(PASTA, 'Relatório Equipamentos - Completo.xls')
_u.ARQ_EQUIP_XLSX = ARQ_EQUIP_XLSX

_lock = threading.Lock()
_seq = [0]

# Equipamento em nome da propria Solivetti nao fica na guia do estado: vai todo
# para a guia "--", seja qual for a UF que vier na planilha. Vale na importacao,
# no cadastro manual e na reconstrucao a partir das planilhas.
UF_PROPRIA = '--'
CLIENTE_PROPRIO = 'SOLIVETTI'


def _e_da_casa(cliente):
    """O equipamento esta em nome da propria empresa?"""
    return CLIENTE_PROPRIO in norm(cliente)


def novo_id(prefixo='x'):
    _seq[0] += 1
    return f'{prefixo}{_seq[0]}'


def construir_seed():
    """Primeira carga: monta o dados.json a partir das planilhas de origem."""
    terc = ler_terceirizados()
    equip = ler_equipamentos()
    tecnicos = {}
    for t in terc:
        tecnicos.setdefault(t['uf'], []).append({
            'id': novo_id('t'),
            'nome': t['nome'], 'local': t['local'], 'valor': t['valor'],
            'telefone': t['telefone'], 'contato': t['contato'],
            'email': t['email'], 'obs': t['obs'],
        })

    equipamentos = []
    for row in equip:
        (serie, modelo, cidade, uf, endereco, bairro,
         complemento, departamento, fabricante, cliente) = row[:10]
        if _e_da_casa(cliente):
            uf = UF_PROPRIA
        equipamentos.append({
            'id': novo_id('e'),
            'serie': serie, 'modelo': modelo, 'cidade': cidade, 'uf': uf,
            'departamento': departamento, 'fabricante': fabricante, 'cliente': cliente,
        })

    return {'tecnicos': tecnicos, 'equipamentos': equipamentos, 'filtros': [], 'chamados': []}


def carregar_dados():
    if os.path.exists(ARQ_DADOS):
        d = None
        ultimo_erro = None
        for _ in range(6):
            try:
                with open(ARQ_DADOS, encoding='utf-8') as f:
                    d = json.load(f)
                break
            except (OSError, ValueError) as e:
                # Outro PC pode estar gravando neste instante: espera e tenta de novo.
                ultimo_erro = e
                time.sleep(0.15)
        if d is None:
            raise ultimo_erro

        maior = 0
        for lst in d.get('tecnicos', {}).values():
            for it in lst:
                _n = _num_id(it.get('id'))
                maior = max(maior, _n)
        for it in d.get('equipamentos', []):
            maior = max(maior, _num_id(it.get('id')))
        for it in d.get('chamados', []):
            maior = max(maior, _num_id(it.get('id')))
        _seq[0] = max(_seq[0], maior)
        d.setdefault('filtros', [])
        # Base feita antes do modulo de chamados existir: abre sem chamado nenhum
        # em vez de quebrar a tela.
        d.setdefault('chamados', [])
        return d
    try:
        d = construir_seed()
    except FileNotFoundError as e:
        # Nao existe dados.json e as planilhas de origem tambem nao estao ai:
        # tipicamente um PC sem o Y: mapeado. Antes isso derrubava o programa, e
        # como o .exe e --noconsole a pessoa nao via nada na tela. Agora abre uma
        # base VAZIA e nao grava nada: se o Y: voltar, o dados.json de verdade
        # ainda sera encontrado na proxima abertura.
        print(f'AVISO: nao achei o dados.json nem as planilhas de origem em:')
        print(f'       {PASTA}')
        print(f'       (faltando: {os.path.basename(e.filename or "")})')
        print(f'AVISO: abrindo uma base VAZIA, visivel so neste PC.')
        return {'tecnicos': {}, 'equipamentos': [], 'filtros': [], 'chamados': []}
    salvar_dados(d)
    return d


def _num_id(s):
    try:
        return int(''.join(ch for ch in str(s) if ch.isdigit()) or 0)
    except Exception:
        return 0


def salvar_dados(d):
    # Grava num temporario e troca de nome: nunca deixa o dados.json pela metade.
    tmp = ARQ_DADOS + f'.{os.getpid()}.tmp'
    with open(tmp, 'w', encoding='utf-8') as f:
        json.dump(d, f, ensure_ascii=False, indent=2)
    os.replace(tmp, ARQ_DADOS)


_ARQ_TRAVA = ARQ_DADOS + '.lock'


class TravaOcupada(Exception):
    """Outro PC está salvando neste instante."""


class trava_rede:
    """Garante que só um PC grave o dados.json por vez (não perde a edição de ninguém)."""

    def __init__(self, timeout=15):
        self.timeout = timeout
        self.fd = None

    def __enter__(self):
        inicio = time.time()
        while True:
            try:
                self.fd = os.open(_ARQ_TRAVA, os.O_CREAT | os.O_EXCL | os.O_RDWR)
                return self
            except FileExistsError:
                try:
                    # Trava velha (PC que travou/desligou no meio): derruba.
                    if time.time() - os.path.getmtime(_ARQ_TRAVA) > 30:
                        os.remove(_ARQ_TRAVA)
                        continue
                except OSError:
                    pass
                if time.time() - inicio > self.timeout:
                    raise TravaOcupada()
                time.sleep(0.15)

    def __exit__(self, *a):
        try:
            if self.fd is not None:
                os.close(self.fd)
        except OSError:
            pass
        try:
            os.remove(_ARQ_TRAVA)
        except OSError:
            pass


DADOS = carregar_dados()


def ip_local():
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.connect(('8.8.8.8', 80))
        ip = s.getsockname()[0]
        s.close()
        return ip
    except Exception:
        return '127.0.0.1'


def _linhas_do_excel(nome, conteudo):
    """Devolve todas as linhas do Excel enviado pelo navegador."""
    ext = nome.lower().rsplit('.', 1)[-1]
    if ext == 'xlsx':
        ws = openpyxl.load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True).active
        return [list(r) for r in ws.iter_rows(values_only=True)]
    sh = xlrd.open_workbook(file_contents=conteudo).sheet_by_index(0)
    return [sh.row_values(i) for i in range(sh.nrows)]


_COLS = {
    'serie': ['SERIE', 'N SERIE', 'NUMERO DE SERIE'],
    'modelo': ['MODELO'],
    'uf': ['UF', 'ESTADO'],
    'cidade': ['CIDADE', 'MUNICIPIO'],
    'departamento': ['DEPARTAMENTO', 'SETOR'],
    'fabricante': ['FABRICANTE', 'MARCA'],
    'cliente': ['CLIENTE', 'RAZAO SOCIAL'],
}


def _mapear_colunas(headers):
    """Descobre em que posicao esta cada coluna, pelo titulo da primeira linha."""
    idx = {}
    for i, h in enumerate(headers):
        hn = norm(h)
        if not hn:
            continue
        for campo, chaves in _COLS.items():
            if campo in idx:
                continue
            if any(k in hn for k in chaves):
                idx[campo] = i
                break
    return idx


# ---------------------------------------------------------------- importacao

def _chave_serie(v):
    return (v or '').strip().upper()


def _mesmo_lugar(atual, novo):
    """A serie continua no mesmo estado, na mesma cidade e no mesmo cliente?"""
    for campo in ('uf', 'cidade', 'cliente'):
        if norm(atual.get(campo, '')) != norm(novo.get(campo, '')):
            return False
    return True


def _lugar(e):
    return {'uf': e.get('uf', ''), 'cidade': e.get('cidade', ''), 'cliente': e.get('cliente', '')}


def _ler_novos(nome, conteudo):
    """Le o Excel e devolve a lista de equipamentos que vieram no arquivo."""
    linhas = _linhas_do_excel(nome, conteudo)
    if not linhas:
        raise ValueError('Arquivo vazio.')
    idx = _mapear_colunas(linhas[0])
    if 'serie' not in idx:
        raise ValueError("Não encontrei a coluna 'Série' na primeira linha do arquivo.")

    novos = []
    for r in linhas[1:]:
        def get(campo):
            i = idx.get(campo)
            if i is not None and i < len(r):
                return limpa(r[i])
            return ''

        serie = get('serie')
        if not serie:
            continue
        cidade = get('cidade')
        uf = get('uf').upper()
        if not uf:
            # Sem coluna UF: tenta achar o "(PE)" no fim do nome da cidade.
            m = re.search(r'\(([A-Za-z]{2})\)', cidade)
            uf = m.group(1).upper() if m else ''
        cliente = get('cliente')
        if _e_da_casa(cliente):
            # Equipamento da propria Solivetti: guia "--", nao a do estado.
            uf = UF_PROPRIA
        novos.append({
            'serie': serie, 'modelo': get('modelo'), 'cidade': cidade, 'uf': uf,
            'departamento': get('departamento'), 'fabricante': get('fabricante'),
            'cliente': cliente,
        })
    return novos, idx


def _aplicar_campos(atual, novo, idx):
    """Copia os campos do arquivo para o registro do painel.

    So sobrescreve com valor de verdade: coluna que nao veio no arquivo, ou que
    veio em branco, nao apaga o que ja estava no painel (preserva edicao manual).
    """
    for campo in ('serie', 'modelo', 'cidade', 'uf', 'departamento', 'fabricante', 'cliente'):
        if campo not in idx:
            continue
        v = novo.get(campo, '')
        if v:
            atual[campo] = v


# Quantos itens detalhados o relatorio devolve para a tela (o total vai sempre).
_MAX_DETALHE = 400


def importar_equipamentos(nome, conteudo, modo='substituir'):
    novos, idx = _ler_novos(nome, conteudo)
    eq = DADOS['equipamentos']

    if modo == 'substituir':
        antigos = len(eq)
        DADOS['equipamentos'] = [{**n, 'id': novo_id('e')} for n in novos]
        return {'modo': 'substituir', 'importados': len(novos), 'removidos': antigos}

    if modo == 'sincronizar':
        return _sincronizar(novos, idx, eq)

    # merge: acrescenta e atualiza, nunca remove
    indice = {_chave_serie(x.get('serie')): x for x in eq}
    add = upd = 0
    for n in novos:
        chave = _chave_serie(n['serie'])
        if chave in indice:
            indice[chave].update(n)
            upd += 1
        else:
            item = {**n, 'id': novo_id('e')}
            eq.append(item)
            indice[chave] = item
            add += 1
    return {'modo': 'merge', 'adicionados': add, 'atualizados': upd}


def _sincronizar(novos, idx, eq):
    """A serie e unica: o arquivo manda onde cada uma esta.

    - mesma serie, mesmo estado/cidade/cliente -> nao mexe
    - mesma serie em outro lugar               -> move (sai de onde estava)
    - serie que nao existe no painel           -> entra
    - serie do painel que nao veio no arquivo  -> sai
    """
    indice = {}
    for x in eq:
        indice.setdefault(_chave_serie(x.get('serie')), x)

    vistas = set()
    inalterados = 0
    movidos = []
    adicionados = []
    duplicadas = 0

    for n in novos:
        chave = _chave_serie(n['serie'])
        if not chave:
            continue
        if chave in vistas:
            # A mesma serie veio duas vezes no arquivo: vale a ultima linha.
            duplicadas += 1
        vistas.add(chave)

        atual = indice.get(chave)
        if atual is None:
            item = {**n, 'id': novo_id('e')}
            eq.append(item)
            indice[chave] = item
            adicionados.append({'serie': item['serie'], **_lugar(item)})
            continue

        if _mesmo_lugar(atual, n):
            # Mesmo lugar: so acerta modelo/fabricante/departamento se mudaram.
            _aplicar_campos(atual, n, idx)
            inalterados += 1
            continue

        de = _lugar(atual)
        _aplicar_campos(atual, n, idx)
        movidos.append({'serie': atual.get('serie', ''), 'de': de, 'para': _lugar(atual)})

    removidos = [{'serie': x.get('serie', ''), **_lugar(x)}
                 for x in eq if _chave_serie(x.get('serie')) not in vistas]
    if removidos:
        eq[:] = [x for x in eq if _chave_serie(x.get('serie')) in vistas]

    return {
        'modo': 'sincronizar',
        'inalterados': inalterados,
        'movidos': movidos[:_MAX_DETALHE],
        'movidos_total': len(movidos),
        'novos': adicionados[:_MAX_DETALHE],
        'novos_total': len(adicionados),
        'removidos': removidos[:_MAX_DETALHE],
        'removidos_total': len(removidos),
        'duplicadas_no_arquivo': duplicadas,
        'total': len(eq),
    }


# ---------------------------------------------------------------- exportacao

def importar_chamados(conteudo, modo='merge'):
    """Importa a aba CHAMADOS TERCEIRIZADOS de um .xlsx para dentro do painel.

    modo 'merge'      : so entra chamado cujo numero de OS ainda nao existe aqui
                        (o padrao -- nunca sobrescreve o que ja foi atualizado
                        a mao no painel).
    modo 'substituir' : apaga todos os chamados e poe os da planilha no lugar.
    """
    tmp = os.path.join(PASTA, f'_chamados_{os.getpid()}.xlsx')
    try:
        with open(tmp, 'wb') as f:
            f.write(conteudo)
        novos, rel = _ch.importar(tmp, DADOS.get('equipamentos', []), novo_id)
    finally:
        try:
            os.remove(tmp)
        except OSError:
            pass

    if modo == 'substituir':
        rel['ignorados'] = 0
        DADOS['chamados'] = novos
        return rel

    atuais = DADOS.setdefault('chamados', [])

    def _chave(c):
        """Como reconhecer que dois registros sao o mesmo chamado.

        Normalmente e o numero da O.S. Mas ha linhas na planilha SEM numero
        nenhum: se a chave fosse so a O.S., o vazio nunca casaria com nada e
        elas entrariam de novo a cada importacao. Nesses casos vale o conjunto
        cliente + parceiro + observacao.
        """
        os_ = str(c.get('os', '')).strip()
        if os_:
            return ('os', os_)
        return ('sem_os', norm(c.get('cliente')), norm(c.get('parceiro')), norm(c.get('obs')))

    ja_tem = {_chave(c) for c in atuais}
    entram = []
    for c in novos:
        k = _chave(c)
        if k in ja_tem:
            continue
        ja_tem.add(k)        # a propria planilha pode repetir a mesma linha
        entram.append(c)
    rel['ignorados'] = len(novos) - len(entram)
    rel['importados'] = len(entram)
    # o relatorio por situacao tem de refletir o que REALMENTE entrou
    rel['por_situacao'] = {}
    for c in entram:
        rel['por_situacao'][c['situacao']] = rel['por_situacao'].get(c['situacao'], 0) + 1
    rel['sem_uf'] = sum(1 for c in entram if not c.get('uf'))
    atuais.extend(entram)
    return rel


def _split_tel(s):
    m = re.match(r'^\((\w+)\)\s*(.*)$', (s or '').strip())
    if m:
        return m.group(1), m.group(2)
    return '', (s or '')


def _salvar_xlsx(caminho, titulo, cabecalho, linhas, com_titulo=False):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = titulo[:31]
    if com_titulo:
        ws.append([titulo])
    ws.append(cabecalho)
    for ln in linhas:
        ws.append(ln)
    tmp = caminho + '.tmp.xlsx'
    wb.save(tmp)
    os.replace(tmp, caminho)


def exportar_planilhas():
    """Grava o conteudo atual do painel de volta nas planilhas .xlsx da rede."""
    linhas_eq = []
    for e in DADOS.get('equipamentos', []):
        linhas_eq.append([
            e.get('serie', ''), e.get('modelo', ''), e.get('cidade', ''), e.get('uf', ''),
            '', '', '', e.get('departamento', ''), e.get('fabricante', ''), e.get('cliente', ''),
        ])
    _salvar_xlsx(ARQ_EQUIP_XLSX, 'Equipamentos',
                 ['Serie', 'Modelo', 'Cidade', 'UF', 'Endereço', 'Bairro',
                  'Complemento', 'Departamento', 'Fabricante', 'Cliente'],
                 linhas_eq)

    linhas_tec = []
    for uf in sorted(DADOS.get('tecnicos', {})):
        for t in DADOS['tecnicos'][uf]:
            ddd, tel = _split_tel(t.get('telefone', ''))
            linhas_tec.append([
                uf, t.get('local', ''), t.get('nome', ''), t.get('contato', ''), t.get('email', ''),
                ddd, tel, t.get('valor', ''), '', '', t.get('obs', ''),
            ])
    _salvar_xlsx(ARQ_TERC, 'EMPRESAS TERCEIRIZADAS',
                 ['UF', 'LOCAL', 'TERCEIRIZADO', 'CONTATO', 'E-MAIL', 'DDD',
                  'TELEFONE / WHATSAPP', 'VALOR', 'RETORNO', 'DESLOCAMENTO', 'OBS'],
                 linhas_tec, com_titulo=True)

    return {'equipamentos': len(linhas_eq), 'tecnicos': len(linhas_tec),
            'arq_eq': os.path.basename(ARQ_EQUIP_XLSX),
            'arq_tec': os.path.basename(ARQ_TERC)}


# ------------------------------------------------------------------ servidor

class Handler(BaseHTTPRequestHandler):
    def log_message(self, *a):
        return

    def _envia(self, corpo, tipo='application/json', status=200):
        if isinstance(corpo, str):
            corpo = corpo.encode('utf-8')
        self.send_response(status)
        self.send_header('Content-Type', f'{tipo}; charset=utf-8')
        self.send_header('Content-Length', str(len(corpo)))
        self.end_headers()
        self.wfile.write(corpo)

    def do_GET(self):
        if self.path in ('/', '/index.html'):
            self._envia(PAGINA, 'text/html')
            return
        if self.path == '/logo.jpg':
            try:
                with open(ARQ_LOGO, 'rb') as f:
                    self._envia(f.read(), 'image/jpeg')
            except Exception:
                self._envia('sem logo', 'text/plain', 404)
            return
        if self.path == '/api/dados':
            with _lock:
                try:
                    globals()['DADOS'] = carregar_dados()
                except Exception:
                    pass
                # A lista de situacoes vai junto para a tela nao ter uma copia
                # propria: a regra mora so no chamados.py.
                resp = dict(DADOS)
                resp['_situacoes'] = [
                    {'chave': k, 'rotulo': v[0], 'cor': v[1], 'aberta': v[2]}
                    for k, v in _ch.SITUACOES.items()
                ]
                self._envia(json.dumps(resp, ensure_ascii=False))
            return
        self._envia('Nao encontrado', 'text/plain', 404)

    def do_POST(self):
        tam = int(self.headers.get('Content-Length', 0))
        body = self.rfile.read(tam).decode('utf-8') if tam else '{}'
        try:
            req = json.loads(body)
        except Exception:
            req = {}

        try:
            with _lock, trava_rede():
                try:
                    globals()['DADOS'] = carregar_dados()
                except Exception:
                    pass

                if self.path == '/api/tecnico':
                    self._op_tecnico(req)
                elif self.path == '/api/equip':
                    self._op_equip(req)
                elif self.path == '/api/restaurar':
                    globals()['DADOS'] = construir_seed()
                elif self.path == '/api/importar':
                    conteudo = base64.b64decode(req.get('dados_b64', ''))
                    stats = importar_equipamentos(req.get('nome', 'arquivo.xlsx'), conteudo,
                                                  req.get('modo', 'substituir'))
                    salvar_dados(DADOS)
                    self._envia(json.dumps({**DADOS, '_import': stats}, ensure_ascii=False))
                    return
                elif self.path == '/api/chamado':
                    self._op_chamado(req)
                elif self.path == '/api/importar_chamados':
                    conteudo = base64.b64decode(req.get('dados_b64', ''))
                    stats = importar_chamados(conteudo, req.get('modo', 'merge'))
                    salvar_dados(DADOS)
                    self._envia(json.dumps({**DADOS, '_impch': stats}, ensure_ascii=False))
                    return
                elif self.path == '/api/filtro':
                    self._op_filtro(req)
                elif self.path == '/api/exportar':
                    info = exportar_planilhas()
                    self._envia(json.dumps({**DADOS, '_export': info}, ensure_ascii=False))
                    return
                else:
                    self._envia('Rota invalida', 'text/plain', 404)
                    return

                salvar_dados(DADOS)
                self._envia(json.dumps(DADOS, ensure_ascii=False))
        except Exception as e:
            self._envia(json.dumps({'erro': str(e)}), status=500)

    def _op_tecnico(self, req):
        op = req.get('op')
        tec = DADOS['tecnicos']
        if op == 'add':
            item, uf = dict(req['item']), req['uf']
            item['id'] = novo_id('t')
            tec.setdefault(uf, []).append(item)
        elif op == 'update':
            item, uf, _id = dict(req['item']), req['uf'], req['id']
            item['id'] = _id
            atual = tec.get(uf, [])
            idx = next((i for i, x in enumerate(atual) if x.get('id') == _id), None)
            if idx is not None:
                atual[idx] = item
                return
            # Mudou de UF: tira da lista antiga e poe na nova.
            for lst in tec.values():
                lst[:] = [x for x in lst if x.get('id') != _id]
            tec.setdefault(uf, []).append(item)
        elif op == 'delete':
            _id = req['id']
            for lst in tec.values():
                lst[:] = [x for x in lst if x.get('id') != _id]
        elif op == 'reorder':
            uf, ordem = req['uf'], req.get('ordem', [])
            lst = tec.get(uf, [])
            pos = {n: i for i, n in enumerate(ordem)}
            lst.sort(key=lambda x: pos.get(x.get('id'), 1000000000))

    def _op_equip(self, req):
        op = req.get('op')
        eq = DADOS['equipamentos']
        if op == 'add':
            item = dict(req['item'])
            if _e_da_casa(item.get('cliente')):
                item['uf'] = UF_PROPRIA
            item['id'] = novo_id('e')
            eq.append(item)
        elif op == 'update':
            item, _id = dict(req['item']), req['id']
            if _e_da_casa(item.get('cliente')):
                item['uf'] = UF_PROPRIA
            item['id'] = _id
            for i, x in enumerate(eq):
                if x.get('id') == _id:
                    eq[i] = item
                    return
        elif op == 'delete':
            _id = req['id']
            eq[:] = [x for x in eq if x.get('id') != _id]

    def _op_chamado(self, req):
        """Cria, altera, reabre e exclui chamado. Toda mudanca entra no historico."""
        op = req.get('op')
        lista = DADOS.setdefault('chamados', [])
        hoje = time.strftime('%Y-%m-%d')

        if op == 'add':
            item = dict(req['item'])
            item['id'] = novo_id('c')
            item.setdefault('situacao', _ch.SIT_PADRAO)
            # Chamado novo ja nasce com data: e justamente o que faltava na planilha.
            if not item.get('aberto_em'):
                item['aberto_em'] = hoje
            item['historico'] = []
            if item['situacao'] not in _ch.ABERTAS:
                item['fechado_em'] = hoje
            _ch.anotar(item, f"Chamado criado como «{_ch.SITUACOES[item['situacao']][0]}»")
            if item.get('obs'):
                _ch.anotar(item, item['obs'])
            lista.append(item)

        elif op == 'update':
            _id = req['id']
            novo = dict(req['item'])
            atual = next((x for x in lista if x.get('id') == _id), None)
            if atual is None:
                return
            sit_antes = atual.get('situacao')
            hist = atual.get('historico', [])
            novo['id'] = _id
            novo['historico'] = hist
            novo.setdefault('aberto_em', atual.get('aberto_em', ''))
            sit_depois = novo.get('situacao', sit_antes)

            if sit_depois != sit_antes:
                de = _ch.SITUACOES.get(sit_antes, (sit_antes,))[0]
                para = _ch.SITUACOES.get(sit_depois, (sit_depois,))[0]
                _ch.anotar(novo, f'{de} → {para}')
                # Fechou agora: marca a data. Reabriu: limpa, senao o contador
                # de dias congelaria na data da conclusao antiga.
                if sit_depois not in _ch.ABERTAS:
                    novo['fechado_em'] = hoje
                else:
                    novo['fechado_em'] = ''

            nota = (req.get('nota') or '').strip()
            if nota:
                _ch.anotar(novo, nota)

            for i, x in enumerate(lista):
                if x.get('id') == _id:
                    lista[i] = novo
                    break

        elif op == 'nota':
            # So acrescenta um andamento, sem mexer no resto.
            _id, texto = req['id'], (req.get('texto') or '').strip()
            ch = next((x for x in lista if x.get('id') == _id), None)
            if ch is not None and texto:
                _ch.anotar(ch, texto)

        elif op == 'delete':
            _id = req['id']
            lista[:] = [x for x in lista if x.get('id') != _id]

    def _op_filtro(self, req):
        op = req.get('op')
        fs = DADOS.setdefault('filtros', [])
        if op == 'add':
            item = dict(req['item'])
            item['id'] = novo_id('f')
            fs.append(item)
        elif op == 'delete':
            _id = req['id']
            fs[:] = [x for x in fs if x.get('id') != _id]


def _log_erro(e):
    import traceback
    try:
        with open(os.path.join(PASTA, 'erro_servidor.txt'), 'w', encoding='utf-8') as f:
            f.write('Erro ao iniciar/rodar o painel:\n\n')
            f.write(''.join(traceback.format_exception(type(e), e, e.__traceback__)))
            f.write("\nDica: talvez o painel ja esteja aberto (porta 8000 em uso). "
                    "Use 'Parar Painel' e tente de novo.")
    except Exception:
        pass


def _ja_esta_rodando():
    c = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
    c.settimeout(0.4)
    try:
        c.connect(('127.0.0.1', PORTA))
        return True
    except OSError:
        return False
    finally:
        c.close()


class ServidorPainel(ThreadingHTTPServer):
    allow_reuse_address = False


def main():
    # PAINEL_SEM_NAVEGADOR=1: sobe o servidor sem abrir janela (teste automatizado).
    _abrir = not os.environ.get('PAINEL_SEM_NAVEGADOR')
    if _ja_esta_rodando():
        try:
            webbrowser.open(f'http://localhost:{PORTA}')
        except Exception:
            pass
        return
    try:
        servidor = ServidorPainel(('127.0.0.1', PORTA), Handler)
    except OSError:
        try:
            webbrowser.open(f'http://localhost:{PORTA}')
        except Exception:
            pass
        return
    except Exception as e:
        _log_erro(e)
        return

    try:
        os.remove(os.path.join(PASTA, 'erro_servidor.txt'))
    except Exception:
        pass

    print('==========================================================')
    print(f'Painel no ar em http://localhost:{PORTA}')
    print('==========================================================')
    try:
        if _abrir:
            webbrowser.open(f'http://localhost:{PORTA}')
    except Exception:
        pass
    try:
        servidor.serve_forever()
    except KeyboardInterrupt:
        print('\nEncerrado.')
    except Exception as e:
        _log_erro(e)


if __name__ == '__main__':
    main()
