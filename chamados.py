# -*- coding: utf-8 -*-
"""Modulo de CHAMADOS terceirizados.

Nasceu para tirar o acompanhamento da planilha do Google ("PLANILHA APOIO -
ATENDIMENTOS TERCEIRIZADOS", aba CHAMADOS TERCEIRIZADOS), onde o status era
texto livre (60 grafias diferentes para 332 chamados) e nao existia data
nenhuma -- nao dava para saber ha quanto tempo um chamado estava parado.

Aqui a situacao e uma lista fixa, a data de abertura existe, e cada alteracao
vira uma linha no historico em vez de apagar a anterior.

Nada neste arquivo toca em tecnicos ou equipamentos.
"""
import re
import unicodedata
from datetime import date, datetime

# ---------------------------------------------------------------- situacoes

# chave -> (rotulo na tela, cor, conta como "em aberto"?)
SITUACOES = {
    'aberto':          ('Aberto',                      '#6b7a90', True),
    'enviado':         ('Enviado ao terceirizado',     '#2f6fd0', True),
    'aguardando_peca': ('Aguardando peça',        '#d08a00', True),
    'peca_enviada':    ('Peça enviada',           '#8a4fd0', True),
    'orcamento':       ('Aguardando aprovação', '#c85a00', True),
    'concluido':       ('Concluído',              '#1f9254', False),
    'cancelado':       ('Cancelado',                   '#99a3b0', False),
}
SIT_PADRAO = 'aberto'
ABERTAS = [k for k, v in SITUACOES.items() if v[2]]


def _nrm(s):
    """Maiusculas, sem acento, sem pontuacao. Para comparar texto baguncado."""
    s = unicodedata.normalize('NFKD', str(s or '')).encode('ascii', 'ignore').decode()
    return ' '.join(re.sub(r'[^A-Z0-9 ]', ' ', s.upper()).split())


# A ordem IMPORTA: vence a primeira que casar.
# Montada em cima dos 60 status reais que existiam na planilha.
_REGRAS = [
    # Antes de tudo: texto que diz que o atendimento NAO resolveu, mesmo
    # contendo a palavra "atendido". Ex.: "ATENDIDO MAS IMPRESSAO CONTINUA EM BRANCO".
    (r'SEM SUCESSO|MAS IMPRESSAO|CONTINUA EM BRANCO', 'aberto'),
    (r'CANCELAD', 'cancelado'),
    # "CONCLUIDO - ENVIAR PECA PARA O TECNICO" e "ENVIAR PECA (CHAMADO CONCLUIDO)"
    # sao conclusoes: por isso concluido vem antes das regras de peca.
    (r'CONCLU|CONLUIDO|FINALIZAD|RESOLVID|ATENDIDO POR|INSTALADO PELO CLIENTE|'
     r'CHAMADO ATENDIDO|REALIZADO A TROCA', 'concluido'),
    (r'ORCAMENTO|APROVACAO', 'orcamento'),
    (r'AGUARDANDO.*PEC|COMPRA DE PECA|CHEGADA DE PEC', 'aguardando_peca'),
    (r'PECA ENVIADA|ENVIAR PECA|PECA NO CLIENTE|PECA PARA', 'peca_enviada'),
    (r'ENVIADO|ENCAMINHADO|WHATSAPP|EMAIL', 'enviado'),
]


def classificar(texto):
    """Texto livre do status antigo -> uma das situacoes fixas."""
    t = _nrm(texto)
    if not t:
        return SIT_PADRAO
    for padrao, sit in _REGRAS:
        if re.search(padrao, t):
            return sit
    return SIT_PADRAO


# ------------------------------------------------------- cidade / UF

# Clientes como "FIORI MARANHAO" ou "LOGPLACE SERGIPE" trazem o ESTADO no nome,
# nao a cidade. Sem isto eles ficariam sem UF nenhuma.
_ESTADOS = {
    'ACRE': 'AC', 'ALAGOAS': 'AL', 'AMAPA': 'AP', 'AMAZONAS': 'AM', 'BAHIA': 'BA',
    'CEARA': 'CE', 'DISTRITO FEDERAL': 'DF', 'ESPIRITO SANTO': 'ES', 'GOIAS': 'GO',
    'MARANHAO': 'MA', 'MATO GROSSO': 'MT', 'MATO GROSSO DO SUL': 'MS',
    'MINAS GERAIS': 'MG', 'PARA': 'PA', 'PARAIBA': 'PB', 'PARANA': 'PR',
    'PERNAMBUCO': 'PE', 'PIAUI': 'PI', 'RIO DE JANEIRO': 'RJ',
    'RIO GRANDE DO NORTE': 'RN', 'RIO GRANDE DO SUL': 'RS', 'RONDONIA': 'RO',
    'RORAIMA': 'RR', 'SANTA CATARINA': 'SC', 'SAO PAULO': 'SP',
    'SERGIPE': 'SE', 'TOCANTINS': 'TO',
}


def indice_cidades(equipamentos):
    """Monta {cidade normalizada: UF} a partir dos equipamentos que o painel ja tem.

    No cadastro de equipamentos a cidade vem como "SALVADOR (BA)" -- o
    parenteses precisa sair, senao nada casa.
    """
    cont = {}
    for e in equipamentos:
        uf = (e.get('uf') or '').strip().upper()
        if not uf or uf == '--':
            continue
        cid = _nrm(re.sub(r'\(.*?\)', ' ', str(e.get('cidade') or '')))
        if len(cid) > 3:
            cont.setdefault(cid, {})
            cont[cid][uf] = cont[cid].get(uf, 0) + 1
    # cidade homonima em 2 estados: fica com o estado que tem mais equipamentos
    return {c: max(ufs.items(), key=lambda x: x[1])[0] for c, ufs in cont.items()}


def deduzir_local(cliente, idx_cidades):
    """Descobre (cidade, UF) a partir do nome do cliente.

    A planilha nunca teve coluna de cidade: ela mora dentro do nome
    ("FIORI SALVADOR", "RCR SAO LUIS"). Devolve ('', '') quando o nome e so
    razao social ("SAO BRAZ S/A") -- ai alguem preenche a mao no painel.
    """
    n = _nrm(cliente)
    if not n:
        return '', ''
    # nome de estado por extenso (mais longo primeiro: "MATO GROSSO DO SUL" antes de "MATO GROSSO")
    for nome in sorted(_ESTADOS, key=len, reverse=True):
        if re.search(r'\b' + re.escape(nome) + r'\b', n):
            return '', _ESTADOS[nome]
    # nome de cidade que o painel ja conhece (mais longa primeiro:
    # "SAO LUIS" antes de "SAO", "JUAZEIRO DO NORTE" antes de "JUAZEIRO")
    for cid in sorted(idx_cidades, key=len, reverse=True):
        if re.search(r'\b' + re.escape(cid) + r'\b', n):
            return cid, idx_cidades[cid]
    return '', ''


# ------------------------------------------------------------- historico

def anotar(ch, texto, quem='painel'):
    """Acrescenta uma linha ao historico do chamado, sem apagar as anteriores."""
    if not texto:
        return
    ch.setdefault('historico', []).append({
        'em': datetime.now().strftime('%Y-%m-%d %H:%M'),
        'quem': quem,
        'texto': str(texto).strip(),
    })


def dias_parado(ch, hoje=None):
    """Dias desde a abertura (ou ate a conclusao). None quando falta a data."""
    ini = (ch.get('aberto_em') or '').strip()
    if not ini:
        return None
    try:
        d0 = datetime.strptime(ini[:10], '%Y-%m-%d').date()
    except ValueError:
        return None
    fim = (ch.get('fechado_em') or '').strip()
    if fim:
        try:
            d1 = datetime.strptime(fim[:10], '%Y-%m-%d').date()
        except ValueError:
            d1 = hoje or date.today()
    else:
        d1 = hoje or date.today()
    return max(0, (d1 - d0).days)


# ----------------------------------------------------------- importacao

_CAB = {
    'os': ('OS', 'O.S', 'CHAMADO', 'NUMERO'),
    'cliente': ('CLIENTE',),
    'status': ('STATUS', 'SITUACAO', 'ANDAMENTO'),
    'parceiro': ('PARCEIRO', 'TERCEIRIZADO', 'EMPRESA'),
    'rastreio': ('RASTREIO', 'RASTREAMENTO', 'CODIGO'),
    'pedido': ('PEDIDO',),
}


def _mapear(headers):
    """Acha em que coluna esta cada campo, pelo nome do cabecalho."""
    idx = {}
    for i, h in enumerate(headers):
        hn = _nrm(h)
        if not hn:
            continue
        for campo, apelidos in _CAB.items():
            if campo in idx:
                continue
            if any(hn == a or hn.startswith(a + ' ') or hn == a.replace('.', '') for a in apelidos):
                idx[campo] = i
    return idx


def _texto(v):
    """Celula -> texto. O Excel entrega numero de OS como 79119.0."""
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, (datetime, date)):
        return v.strftime('%Y-%m-%d')
    return str(v).strip()


def ler_planilha(caminho, aba='CHAMADOS TERCEIRIZADOS'):
    """Le a aba de chamados do .xlsx e devolve uma lista de dicionarios crus."""
    import openpyxl
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    if aba not in wb.sheetnames:
        raise ValueError(f'A planilha nao tem a aba "{aba}". Abas: {wb.sheetnames}')
    ws = wb[aba]

    linhas = [list(r) for r in ws.iter_rows(values_only=True)]
    # A linha 1 e o titulo com a logo; o cabecalho de verdade e a primeira
    # linha que tiver "OS" e "CLIENTE".
    pos_cab = None
    for i, linha in enumerate(linhas[:10]):
        vals = [_nrm(c) for c in linha]
        if 'OS' in vals and 'CLIENTE' in vals:
            pos_cab = i
            break
    if pos_cab is None:
        raise ValueError('Nao achei o cabecalho (uma linha com "OS" e "CLIENTE").')

    idx = _mapear(linhas[pos_cab])
    faltando = [c for c in ('os', 'cliente') if c not in idx]
    if faltando:
        raise ValueError(f'Faltam colunas na planilha: {", ".join(faltando)}')

    out = []
    for linha in linhas[pos_cab + 1:]:
        reg = {campo: _texto(linha[i]) if i < len(linha) else ''
               for campo, i in idx.items()}
        if not any(reg.values()):
            continue
        out.append(reg)
    return out


def importar(caminho, equipamentos, novo_id, aba='CHAMADOS TERCEIRIZADOS'):
    """Planilha -> lista de chamados prontos para entrar no dados.json.

    `novo_id` e a funcao de id do servidor, para os ids seguirem a mesma serie.
    Devolve (chamados, relatorio).
    """
    crus = ler_planilha(caminho, aba)
    idx_cid = indice_cidades(equipamentos)

    # A mesma O.S. aparece mais de uma vez na planilha: como la nao dava para
    # atualizar um chamado sem apagar o texto anterior, as pessoas lancavam uma
    # linha nova. Aqui as linhas da mesma O.S. viram UM chamado -- a ultima
    # manda nos campos (e a mais recente) e as anteriores entram no historico.
    grupos = {}
    for i, c in enumerate(crus):
        os_ = str(c.get('os', '')).strip()
        chave = ('os', os_) if os_ else ('linha', i)
        grupos.setdefault(chave, []).append(c)

    chamados = []
    rel = {'lidos': len(crus), 'por_situacao': {}, 'sem_uf': 0,
           'sem_data': 0, 'fundidos': 0}
    for linhas in grupos.values():
        c = linhas[-1]
        cidade, uf = deduzir_local(c.get('cliente'), idx_cid)
        status_antigo = c.get('status', '')
        sit = classificar(status_antigo)

        ch = {
            'id': novo_id('c'),
            'os': c.get('os', ''),
            'cliente': c.get('cliente', ''),
            'cidade': cidade,
            'uf': uf,
            'parceiro': c.get('parceiro', ''),
            'situacao': sit,
            'rastreio': c.get('rastreio', ''),
            'pedido': c.get('pedido', ''),
            # A planilha de origem NAO tinha data nenhuma. Nao da para inventar:
            # fica vazio e o painel mostra "sem data" ate alguem preencher.
            'aberto_em': '',
            'fechado_em': '',
            'obs': status_antigo,
            'historico': [],
        }
        # Historico em ordem: as linhas antigas primeiro, a atual por ultimo.
        for ant in linhas[:-1]:
            txt = (ant.get('status') or '').strip()
            quem_ant = (ant.get('parceiro') or '').strip()
            if txt or quem_ant:
                anotar(ch, f'(lancamento anterior na planilha{" - " + quem_ant if quem_ant else ""}) '
                           f'{txt or "sem status"}', quem='planilha')
        if status_antigo:
            anotar(ch, status_antigo, quem='planilha')
        if len(linhas) > 1:
            rel['fundidos'] += len(linhas) - 1
        chamados.append(ch)

        rel['por_situacao'][sit] = rel['por_situacao'].get(sit, 0) + 1
        if not uf:
            rel['sem_uf'] += 1
        rel['sem_data'] += 1

    rel['importados'] = len(chamados)
    return chamados, rel
